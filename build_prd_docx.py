from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "TransitHub_PRD.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True)
        set_cell_shading(table.rows[0].cells[index], "E8EEF5")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], str(value))
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    doc.add_paragraph()
    return table


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def para(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(11)
for style_name, size, color in [
    ("Heading 1", 16, RGBColor(46, 116, 181)),
    ("Heading 2", 13, RGBColor(46, 116, 181)),
    ("Heading 3", 12, RGBColor(31, 77, 120)),
]:
    style = styles[style_name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = color

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("TransitHub PRD and File Documentation")
run.bold = True
run.font.size = Pt(22)
run.font.color.rgb = RGBColor(11, 37, 69)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Responsive Smart Travel Booking and Reservation System")
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(85, 85, 85)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("Prepared for project explanation, submission, and viva reference").italic = True
doc.add_paragraph()

heading(doc, "1. Product Overview")
para(doc, "TransitHub is a responsive smart travel booking platform inspired by RedBus, IRCTC, and MakeMyTrip. It allows users to search bus and train routes, compare fares, select seats, enter passenger details, simulate payment, generate a ticket, view dashboard history, and interact with a rule-based AI travel assistant.")
add_table(doc, ["Field", "Details"], [
    ["Project Name", "TransitHub"],
    ["Project Type", "Responsive Smart Travel Booking Platform"],
    ["Frontend Stack", "HTML5, CSS3, Bootstrap 5, Bootstrap Icons, JavaScript ES6, jQuery"],
    ["Data Layer", "JSON route data and browser LocalStorage"],
    ["Utility Scripts", "Python scripts for generating larger route datasets"],
    ["Current Dataset", "routes.json currently contains 811 route records"],
], [1.7, 4.7])

heading(doc, "2. Product Goals")
add_bullets(doc, [
    "Provide a clean and modern booking experience for bus and train travel.",
    "Allow route search by source, destination, date, and travel type.",
    "Render route results dynamically from JSON data.",
    "Support route filtering, sorting, seat selection, payment simulation, and ticket generation.",
    "Store user bookings in LocalStorage for dashboard history.",
    "Demonstrate syllabus concepts such as DOM manipulation, events, arrays, objects, validation, JSON, LocalStorage, Bootstrap, and jQuery.",
])

heading(doc, "3. User Flow")
add_numbered(doc, [
    "User opens index.html and enters trip search details.",
    "Search details are stored in LocalStorage and the user moves to search.html.",
    "User filters/sorts available route cards and selects a route.",
    "Selected route is saved and seats.html displays an interactive seat grid.",
    "User selects seats and continues to booking.html.",
    "Passenger forms are generated based on selected seats.",
    "User enters passenger details and chooses a simulated payment option.",
    "A digital ticket is generated and saved in booking history.",
    "dashboard.html displays saved bookings and supports cancellation.",
    "chatbot.html answers common travel questions through keyword matching.",
])

heading(doc, "4. Functional Requirements")
add_table(doc, ["Module", "Requirement"], [
    ["Home", "Hero, search form, feature cards, popular destinations, counters, testimonials, footer."],
    ["Search", "Dynamic route cards, filters, price range, sorting, and route selection."],
    ["Seats", "Clickable grid, booked seats disabled, selected seats highlighted, fare updates."],
    ["Booking", "Passenger forms per selected seat, validation, payment simulation, ticket generation."],
    ["Dashboard", "Booking history, profile card, saved destinations, cancellation, clearing history."],
    ["Chatbot", "Rule-based travel assistant with quick prompts and keyword responses."],
], [1.4, 5.0])

heading(doc, "5. Architecture")
para(doc, "TransitHub uses a static frontend architecture. HTML pages define structure, CSS files define presentation, JavaScript modules provide behavior, JSON files provide route data, and LocalStorage stores temporary user state.")
add_table(doc, ["Layer", "Responsibility"], [
    ["User Interface", "HTML pages and Bootstrap components."],
    ["Presentation", "style.css, animations.css, responsive.css."],
    ["Logic", "JavaScript modules in the js folder."],
    ["Data", "routes.json and trains_dict.json."],
    ["Persistence", "Browser LocalStorage."],
], [1.5, 4.9])

heading(doc, "6. File-by-File Documentation")

file_rows = [
    ["index.html", "Landing page with navbar, hero search, features, popular destinations, counters, testimonials, footer, and shared scripts."],
    ["search.html", "Route results page with search form, filter sidebar, sorting controls, results summary, and dynamic results list."],
    ["seats.html", "Interactive seat selection page with seat legend, grid container, selected seat summary, and fare total."],
    ["booking.html", "Passenger details, payment simulation, booking summary, digital ticket, and ticket print/download UI."],
    ["dashboard.html", "User dashboard with profile card, saved destinations, booking history, cancellation, and clear history."],
    ["chatbot.html", "AI Travel Assistant page with chat messages, quick prompts, input box, and send button."],
    ["css/style.css", "Main design system: theme variables, navbar, hero, cards, forms, route cards, seats, ticket, dashboard, chatbot, and buttons."],
    ["css/animations.css", "Animation keyframes and reveal classes for hero, cards, buttons, seat selection, shimmer, and reduced-motion support."],
    ["css/responsive.css", "Media queries for mobile/tablet layouts, stacked cards, smaller seats, and payment/chat responsiveness."],
    ["js/storage.js", "LocalStorage helper for safe JSON get, set, and remove operations."],
    ["js/app.js", "Shared app utilities: load routes, theme toggle, toasts, duration formatting, page motion, hero words, homepage search, counters, popular routes."],
    ["js/search.js", "Loads routes, filters/sorts results, renders route cards, stores selected route, and redirects to seat selection."],
    ["js/seats.js", "Builds seat grid, manages selected/booked seats, calculates fare, stores seats, and guards booking continuation."],
    ["js/booking.js", "Builds passenger forms per seat, validates inputs, creates booking object, saves history, and renders ticket."],
    ["js/dashboard.js", "Reads stored bookings, renders booking history, cancels tickets, and clears history."],
    ["js/chatbot.js", "Rule-based assistant using keyword arrays and predefined responses."],
    ["data/routes.json", "Main route database consumed by search and homepage popular route rendering."],
    ["data/trains_dict.json", "Raw train dataset used by convert.py to generate train routes."],
    ["convert.py", "Converts raw train records into TransitHub route objects and appends them to routes.json."],
    ["generate_buses.py", "Generates random bus route records for popular city pairs and appends them to routes.json."],
    ["assets/", "Reserved for local images, videos, icons, and future static media."],
    ["bootstrap/", "Reserved for local Bootstrap assets if the project is made offline."],
    ["TransitHub.zip", "Compressed project package for sharing or submission."],
]
add_table(doc, ["File/Folder", "Purpose"], file_rows, [1.55, 4.85])

heading(doc, "7. Data Model")
heading(doc, "7.1 Route Object", 2)
add_table(doc, ["Property", "Meaning"], [
    ["id", "Unique route identifier."],
    ["operator", "Bus company or train name."],
    ["type", "Bus or Train."],
    ["from / to", "Source and destination city."],
    ["departure / arrival", "Displayed travel timing."],
    ["duration", "Trip duration in minutes."],
    ["seatsLeft", "Available seats shown on route card."],
    ["rating", "Customer rating used for sorting."],
    ["fare", "Base fare per selected seat."],
    ["amenities", "Feature tags such as AC, Non-AC, Sleeper."],
    ["time", "Morning or Evening filter value."],
    ["image", "Route/destination image URL."],
], [1.5, 4.9])

heading(doc, "7.2 Booking Object", 2)
para(doc, "A booking stores ticket ID, passenger array, selected route, selected seats, payment method, fare, and date. It is saved inside the bookings LocalStorage key.")

heading(doc, "8. LocalStorage Keys")
add_table(doc, ["Key", "Purpose"], [
    ["theme", "Stores dark/light mode preference."],
    ["searchQuery", "Stores source, destination, date, passengers, and travel type."],
    ["selectedRoute", "Stores route selected from search results."],
    ["selectedSeats", "Stores chosen seat numbers."],
    ["bookings", "Stores generated ticket history."],
    ["lastTicket", "Stores the most recently generated ticket."],
], [1.6, 4.8])

heading(doc, "9. JavaScript Concepts Demonstrated")
add_bullets(doc, [
    "DOM manipulation and dynamic HTML rendering.",
    "Event handling for forms, clicks, filters, seat selection, and chatbot messages.",
    "Arrays and objects for route, passenger, seat, and booking data.",
    "JSON parsing through fetch and route dataset loading.",
    "LocalStorage for browser-side persistence.",
    "Regex validation for name, email, and phone inputs.",
    "Filtering and sorting route arrays.",
    "jQuery DOM updates and counter animation.",
    "Bootstrap components such as navbar, carousel, forms, buttons, badges, and toasts.",
])

heading(doc, "10. External Libraries")
add_table(doc, ["Library", "Use"], [
    ["Bootstrap 5", "Responsive layout, navbar, grid, forms, buttons, carousel, utility classes."],
    ["Bootstrap Icons", "Icons in navbar, feature cards, search buttons, and ticket/download actions."],
    ["jQuery", "DOM manipulation and animated counters."],
    ["Unsplash URLs", "Travel-related image backgrounds and route images."],
], [1.5, 4.9])

heading(doc, "11. Current Limitations")
add_bullets(doc, [
    "No backend server or centralized database.",
    "No real authentication or login system.",
    "No real payment gateway.",
    "Seat booking state is simulated and not synchronized across devices.",
    "Chatbot is rule-based and not connected to a real AI model.",
    "LocalStorage data can be cleared by the browser user.",
])

heading(doc, "12. Future Scope")
add_bullets(doc, [
    "Add backend using Flask, Django, Node.js, or Express.",
    "Use MongoDB, MySQL, PostgreSQL, or Firebase for persistent storage.",
    "Add real user login and admin dashboard.",
    "Integrate payment gateway simulation or real payment APIs.",
    "Generate actual QR codes for tickets.",
    "Send ticket confirmation by email or SMS.",
    "Add real AI assistant integration using Gemini or OpenAI API.",
    "Add live route availability, cancellation policies, and PNR-style tracking.",
])

heading(doc, "13. Viva Explanation Summary")
para(doc, "TransitHub is a responsive travel booking website built with HTML, CSS, Bootstrap, JavaScript, jQuery, JSON, and LocalStorage. It supports searching buses and trains, filtering and sorting results, selecting seats, entering passenger details, simulating payment, generating tickets, and managing booking history. The route database is stored in JSON and can be expanded using Python scripts. The AI Travel Assistant is implemented as a rule-based chatbot using JavaScript keyword matching.")

doc.save(OUT)
print(OUT)
